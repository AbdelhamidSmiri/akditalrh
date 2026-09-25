"""Imports Excel via openpyxl."""
from datetime import datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


def _cell(row, idx, default=None):
    if idx >= len(row):
        return default
    val = row[idx]
    return default if val is None else val


def _str(val):
    if val is None:
        return None
    return str(val).strip()


def _int(val):
    if val is None or val == "":
        return None
    try:
        return int(float(val))
    except (TypeError, ValueError):
        return None


def _decimal(val):
    if val is None or val == "":
        return None
    try:
        return Decimal(str(val))
    except (InvalidOperation, ValueError):
        return None


def _date(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if hasattr(val, "year"):
        return val
    try:
        return datetime.strptime(str(val)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def iter_rows(file_storage, min_row=2):
    """Itère les lignes d'un classeur Excel (1ère feuille)."""
    wb = load_workbook(file_storage, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        yield list(row)
    wb.close()


def parse_voitures(file_storage) -> list[dict]:
    """Colonnes attendues : user_id, ste_location, immatriculation, marque, modele, ..."""
    rows = []
    for row in iter_rows(file_storage):
        rows.append({
            "user_id": _int(_cell(row, 0)),
            "ste_location": _str(_cell(row, 1)),
            "immatriculation": _str(_cell(row, 2)),
            "marque": _str(_cell(row, 3)),
            "modele": _str(_cell(row, 4)),
            "valeur_locative_ht": _decimal(_cell(row, 5)),
            "duree_mois": _int(_cell(row, 6)),
            "km": _int(_cell(row, 7)),
            "cout_km_suplm_ttc": _decimal(_cell(row, 8)),
            "nbr_pneus": _int(_cell(row, 9)),
            "service_link": _str(_cell(row, 10)),
            "categorie_vr": _str(_cell(row, 11)),
            "livraison": _date(_cell(row, 12)),
            "fin_contrat": _date(_cell(row, 13)),
            "km_actuel": _int(_cell(row, 14)),
        })
    return rows


def parse_referentiels(file_storage) -> list[dict]:
    rows = []
    for row in iter_rows(file_storage):
        rows.append({
            "modele": _str(_cell(row, 0)) or "",
            "duree": _str(_cell(row, 1)) or "",
            "kilometrage": _str(_cell(row, 2)) or "",
            "loyer_ht_ald": _decimal(_cell(row, 3)),
            "loyer_ttc_ald": _decimal(_cell(row, 4)),
            "km_supp_ht_ald": _decimal(_cell(row, 5)),
            "nb_pneus_ald": _int(_cell(row, 6)),
            "cat_remplacement_ald": _str(_cell(row, 7)),
            "loyer_ht_arval": _decimal(_cell(row, 8)),
            "loyer_ttc_arval": _decimal(_cell(row, 9)),
            "km_supp_ht_arval": _decimal(_cell(row, 10)),
            "nb_pneus_arval": _int(_cell(row, 11)),
            "cat_remplacement_arval": _str(_cell(row, 12)),
            "moyenne_offres_ttc": _decimal(_cell(row, 13)),
            "valeur_actuelle_moy_ttc": _decimal(_cell(row, 14)),
            "augmentation_pourcent": _str(_cell(row, 15)),
            "categorie": _str(_cell(row, 16)),
        })
    return rows


def parse_cartes_carburant(file_storage) -> list[dict]:
    rows = []
    for row in iter_rows(file_storage):
        rows.append({
            "user_id": _int(_cell(row, 0)) or 0,
            "site_id": _int(_cell(row, 1)) or 0,
            "prestataire": _str(_cell(row, 2)) or "AFRIQUIA",
            "identifiant": _str(_cell(row, 3)) or "",
            "carte_labelle": _str(_cell(row, 4)) or "",
            "utilisateur": _str(_cell(row, 5)) or "",
            "nv_plafond": _decimal(_cell(row, 6)),
            "type_mt": _str(_cell(row, 7)),
            "status": _str(_cell(row, 8)) or "Active",
        })
    return rows


def parse_tag_jawaz(file_storage) -> list[dict]:
    rows = []
    for row in iter_rows(file_storage):
        rows.append({
            "user_id": _int(_cell(row, 0)) or 0,
            "site_id": _int(_cell(row, 1)) or 0,
            "ref_jawaz": _str(_cell(row, 2)) or "",
        })
    return rows
